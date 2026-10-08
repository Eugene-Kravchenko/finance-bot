import logging
import os
import sqlite3
import tempfile
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    ConversationHandler,
    MessageHandler,
    filters,
)

logging.basicConfig(
    format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    level=logging.INFO,
)
LOG = logging.getLogger(__name__)

TZ_NAME = os.getenv("TZ", "Europe/Moscow")
TZ = ZoneInfo(TZ_NAME)
DB_PATH = Path(os.getenv("DB_PATH", "data/mood.sqlite3"))
CHECKIN_TIMES = [
    item.strip() for item in os.getenv("CHECKIN_TIMES", "10:00,16:00,22:00").split(",")
    if item.strip()
]
ADMIN_CHAT_ID = int(os.getenv("ADMIN_CHAT_ID", "0") or 0)

CONSENT, MOOD, ENERGY, ANXIETY, SLEEP, MEDS, SIDE_EFFECTS, NOTE, SAFETY = range(9)


def db() -> sqlite3.Connection:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db() -> None:
    with db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS participants (
                chat_id INTEGER PRIMARY KEY,
                first_name TEXT,
                username TEXT,
                consented_at TEXT NOT NULL,
                active INTEGER NOT NULL DEFAULT 1
            );

            CREATE TABLE IF NOT EXISTS checkins (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                chat_id INTEGER NOT NULL,
                started_at TEXT NOT NULL,
                completed_at TEXT NOT NULL,
                period TEXT NOT NULL,
                mood INTEGER NOT NULL,
                energy INTEGER NOT NULL,
                anxiety INTEGER NOT NULL,
                sleep_hours REAL,
                meds TEXT NOT NULL,
                side_effects TEXT,
                note TEXT,
                safety TEXT,
                FOREIGN KEY(chat_id) REFERENCES participants(chat_id)
            );
            CREATE INDEX IF NOT EXISTS idx_checkins_chat_time
            ON checkins(chat_id, completed_at);
            """
        )


def scale_keyboard(prefix: str, minimum: int = 1, maximum: int = 10) -> InlineKeyboardMarkup:
    rows = []
    values = list(range(minimum, maximum + 1))
    for start in range(0, len(values), 5):
        rows.append([
            InlineKeyboardButton(str(value), callback_data=f"{prefix}:{value}")
            for value in values[start:start + 5]
        ])
    return InlineKeyboardMarkup(rows)


def choice_keyboard(prefix: str, choices: list[tuple[str, str]]) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup([
        [InlineKeyboardButton(label, callback_data=f"{prefix}:{value}")]
        for label, value in choices
    ])


def period_now() -> str:
    hour = datetime.now(TZ).hour
    if hour < 13:
        return "утро"
    if hour < 19:
        return "день"
    return "вечер"


def participant_is_active(chat_id: int) -> bool:
    with db() as conn:
        row = conn.execute(
            "SELECT active FROM participants WHERE chat_id = ?", (chat_id,)
        ).fetchone()
    return bool(row and row["active"])


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    assert update.effective_chat and update.effective_user and update.message
    chat_id = update.effective_chat.id
    if participant_is_active(chat_id):
        await update.message.reply_text(
            "Бот уже подключён. Я буду ненавязчиво спрашивать о состоянии "
            f"в {', '.join(CHECKIN_TIMES)} ({TZ_NAME}).\n\n"
            "Команды: /checkin, /report, /export, /pause, /resume, /delete_my_data"
        )
        return ConversationHandler.END

    keyboard = InlineKeyboardMarkup([[
        InlineKeyboardButton("Согласна", callback_data="consent:yes"),
        InlineKeyboardButton("Не согласна", callback_data="consent:no"),
    ]])
    await update.message.reply_text(
        "Этот бот помогает отмечать настроение и самочувствие. Ответы сохраняются "
        "в закрытой базе и могут быть выгружены в Excel. Если задан администратор, "
        "он также сможет получить сводку. Это не медицинская диагностика и не замена врачу.\n\n"
        "Ты согласна на хранение этих данных?",
        reply_markup=keyboard,
    )
    return CONSENT


async def consent(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    query = update.callback_query
    assert query and query.message and update.effective_user and update.effective_chat
    await query.answer()
    if query.data != "consent:yes":
        await query.edit_message_text("Хорошо. Данные не сохраняются, напоминаний не будет.")
        return ConversationHandler.END

    user = update.effective_user
    with db() as conn:
        conn.execute(
            """INSERT INTO participants(chat_id, first_name, username, consented_at, active)
               VALUES (?, ?, ?, ?, 1)
               ON CONFLICT(chat_id) DO UPDATE SET
                 first_name=excluded.first_name,
                 username=excluded.username,
                 consented_at=excluded.consented_at,
                 active=1""",
            (
                update.effective_chat.id,
                user.first_name,
                user.username,
                datetime.now(TZ).isoformat(),
            ),
        )
    await query.edit_message_text(
        "Готово. Нажми /checkin, чтобы сделать первую отметку. "
        "В любой момент можно приостановить напоминания командой /pause."
    )
    return ConversationHandler.END


async def begin_checkin(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    chat = update.effective_chat
    if not chat or not participant_is_active(chat.id):
        if update.message:
            await update.message.reply_text("Сначала нажми /start и подтверди согласие.")
        return ConversationHandler.END

    context.user_data.clear()
    context.user_data["started_at"] = datetime.now(TZ).isoformat()
    context.user_data["period"] = period_now()
    message = update.message or (update.callback_query and update.callback_query.message)
    assert message
    await message.reply_text(
        "Как настроение сейчас?\n1 — очень тяжело, 10 — очень хорошо.",
        reply_markup=scale_keyboard("mood"),
    )
    return MOOD


async def scheduled_prompt(context: ContextTypes.DEFAULT_TYPE) -> None:
    with db() as conn:
        rows = conn.execute("SELECT chat_id FROM participants WHERE active = 1").fetchall()
    keyboard = InlineKeyboardMarkup([[
        InlineKeyboardButton("Отметить состояние", callback_data="begin:scheduled")
    ]])
    for row in rows:
        try:
            await context.bot.send_message(
                row["chat_id"],
                "Короткая отметка самочувствия — около минуты.",
                reply_markup=keyboard,
            )
        except Exception:
            LOG.exception("Не удалось отправить напоминание chat_id=%s", row["chat_id"])


async def begin_button(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    query = update.callback_query
    assert query
    await query.answer()
    return await begin_checkin(update, context)


async def numeric_answer(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    query = update.callback_query
    assert query and query.data and query.message
    await query.answer()
    key, raw_value = query.data.split(":", 1)
    context.user_data[key] = int(raw_value)

    if key == "mood":
        await query.edit_message_text(f"Настроение: {raw_value}/10")
        await query.message.reply_text("Сколько энергии?", reply_markup=scale_keyboard("energy"))
        return ENERGY
    if key == "energy":
        await query.edit_message_text(f"Энергия: {raw_value}/10")
        await query.message.reply_text(
            "Насколько сильная тревога?\n1 — почти нет, 10 — максимальная.",
            reply_markup=scale_keyboard("anxiety"),
        )
        return ANXIETY

    await query.edit_message_text(f"Тревога: {raw_value}/10")
    await query.message.reply_text(
        "Сколько часов спала за последние сутки? Напиши число, например: 7.5\n"
        "Если не хочешь отвечать — /skip"
    )
    return SLEEP


async def sleep_answer(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    assert update.message
    text = update.message.text.replace(",", ".").strip()
    try:
        value = float(text)
        if not 0 <= value <= 24:
            raise ValueError
    except ValueError:
        await update.message.reply_text("Нужно число от 0 до 24, например 7.5. Или /skip.")
        return SLEEP
    context.user_data["sleep_hours"] = value
    return await ask_meds(update)


async def skip_sleep(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    context.user_data["sleep_hours"] = None
    return await ask_meds(update)


async def ask_meds(update: Update) -> int:
    assert update.message
    await update.message.reply_text(
        "Лекарства сегодня приняты по назначению врача?",
        reply_markup=choice_keyboard("meds", [
            ("Да", "yes"), ("Частично", "partial"), ("Нет", "no"), ("Не применимо", "na")
        ]),
    )
    return MEDS


async def meds_answer(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    query = update.callback_query
    assert query and query.data and query.message
    await query.answer()
    context.user_data["meds"] = query.data.split(":", 1)[1]
    await query.edit_message_text("Ответ о лекарствах сохранён.")
    await query.message.reply_text(
        "Есть ли заметные побочные эффекты или необычные изменения? "
        "Напиши коротко или /skip."
    )
    return SIDE_EFFECTS


async def side_effects_answer(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    assert update.message
    context.user_data["side_effects"] = update.message.text.strip()
    return await ask_note(update)


async def skip_side_effects(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    context.user_data["side_effects"] = ""
    return await ask_note(update)


async def ask_note(update: Update) -> int:
    assert update.message
    await update.message.reply_text(
        "Что сильнее всего повлияло на состояние? Можно написать одной фразой или /skip."
    )
    return NOTE


async def note_answer(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    assert update.message
    context.user_data["note"] = update.message.text.strip()
    return await ask_safety(update, context)


async def skip_note(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    context.user_data["note"] = ""
    return await ask_safety(update, context)


async def ask_safety(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    assert update.message
    if context.user_data.get("mood", 10) <= 2:
        await update.message.reply_text(
            "Есть ли прямо сейчас мысли причинить себе вред или ощущение, что ты не в безопасности?",
            reply_markup=choice_keyboard("safety", [("Нет", "no"), ("Да", "yes")]),
        )
        return SAFETY
    context.user_data["safety"] = "not_asked"
    return await finish_checkin(update, context)


async def safety_answer(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    query = update.callback_query
    assert query and query.data and query.message
    await query.answer()
    answer = query.data.split(":", 1)[1]
    context.user_data["safety"] = answer
    await query.edit_message_text("Ответ сохранён.")
    if answer == "yes":
        await query.message.reply_text(
            "Пожалуйста, не оставайся с этим одна. Свяжись сейчас с человеком, которому доверяешь, "
            "или с местной экстренной службой. Если есть непосредственная опасность — звони 112. "
            "Бот не может обеспечить экстренную помощь."
        )
    return await finish_checkin(update, context, message=query.message)


async def finish_checkin(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
    message=None,
) -> int:
    chat = update.effective_chat
    assert chat
    completed_at = datetime.now(TZ).isoformat()
    with db() as conn:
        conn.execute(
            """INSERT INTO checkins(
                 chat_id, started_at, completed_at, period, mood, energy, anxiety,
                 sleep_hours, meds, side_effects, note, safety
               ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                chat.id,
                context.user_data["started_at"],
                completed_at,
                context.user_data["period"],
                context.user_data["mood"],
                context.user_data["energy"],
                context.user_data["anxiety"],
                context.user_data.get("sleep_hours"),
                context.user_data["meds"],
                context.user_data.get("side_effects", ""),
                context.user_data.get("note", ""),
                context.user_data.get("safety", "not_asked"),
            ),
        )
    target = message or update.message
    assert target
    await target.reply_text("Готово. Спасибо — отметка сохранена.")
    context.user_data.clear()
    return ConversationHandler.END


def report_text(chat_id: int, days: int = 7) -> str:
    since = (datetime.now(TZ) - timedelta(days=days)).isoformat()
    with db() as conn:
        rows = conn.execute(
            """SELECT mood, energy, anxiety, sleep_hours, meds
               FROM checkins WHERE chat_id = ? AND completed_at >= ?
               ORDER BY completed_at""",
            (chat_id, since),
        ).fetchall()
    if not rows:
        return "За последние 7 дней отметок пока нет."

    def avg(field: str) -> str:
        values = [float(row[field]) for row in rows if row[field] is not None]
        return f"{sum(values) / len(values):.1f}" if values else "—"

    meds_ok = sum(row["meds"] == "yes" for row in rows)
    return (
        f"Сводка за 7 дней ({len(rows)} отметок):\n"
        f"• настроение: {avg('mood')}/10\n"
        f"• энергия: {avg('energy')}/10\n"
        f"• тревога: {avg('anxiety')}/10\n"
        f"• сон: {avg('sleep_hours')} ч\n"
        f"• лекарства по назначению: {meds_ok} из {len(rows)} отметок\n\n"
        "Это наблюдение за динамикой, а не медицинское заключение."
    )


async def report(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    assert update.effective_chat and update.message
    target = report_target(update.effective_chat.id)
    if target is None:
        await update.message.reply_text("Пока нет подключённой участницы.")
        return
    await update.message.reply_text(report_text(target))


def report_target(requester_chat_id: int) -> int | None:
    if not ADMIN_CHAT_ID or requester_chat_id != ADMIN_CHAT_ID:
        return requester_chat_id
    with db() as conn:
        row = conn.execute(
            """SELECT chat_id FROM participants
               WHERE chat_id != ? ORDER BY consented_at DESC LIMIT 1""",
            (ADMIN_CHAT_ID,),
        ).fetchone()
    return int(row["chat_id"]) if row else None


def create_xlsx(chat_id: int, output: Path) -> None:
    with db() as conn:
        rows = conn.execute(
            "SELECT * FROM checkins WHERE chat_id = ? ORDER BY completed_at", (chat_id,)
        ).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Наблюдения"
    headers = [
        "Дата и время", "Период", "Настроение", "Энергия", "Тревога",
        "Сон, ч", "Лекарства", "Побочные эффекты", "Комментарий", "Безопасность",
    ]
    ws.append(headers)
    for row in rows:
        dt = datetime.fromisoformat(row["completed_at"])
        ws.append([
            dt.strftime("%d.%m.%Y %H:%M"), row["period"], row["mood"], row["energy"],
            row["anxiety"], row["sleep_hours"], row["meds"], row["side_effects"],
            row["note"], row["safety"],
        ])
    fill = PatternFill("solid", fgColor="D9EAD3")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    widths = [19, 12, 14, 11, 11, 10, 16, 28, 38, 16]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(output)


async def export(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    assert update.effective_chat and update.message
    target = report_target(update.effective_chat.id)
    if target is None:
        await update.message.reply_text("Пока нет подключённой участницы.")
        return
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "mood_report.xlsx"
        create_xlsx(target, path)
        with path.open("rb") as file:
            await update.message.reply_document(file, filename="mood_report.xlsx")


async def my_id(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    assert update.effective_chat and update.message
    await update.message.reply_text(f"ID этого чата: {update.effective_chat.id}")


async def pause(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    assert update.effective_chat and update.message
    with db() as conn:
        conn.execute("UPDATE participants SET active = 0 WHERE chat_id = ?", (update.effective_chat.id,))
    await update.message.reply_text("Напоминания приостановлены. Вернуть их: /resume")


async def resume(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    assert update.effective_chat and update.message
    with db() as conn:
        conn.execute("UPDATE participants SET active = 1 WHERE chat_id = ?", (update.effective_chat.id,))
    await update.message.reply_text("Напоминания снова включены.")


async def delete_my_data(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    assert update.effective_chat and update.message
    keyboard = InlineKeyboardMarkup([[
        InlineKeyboardButton("Да, удалить", callback_data="delete:yes"),
        InlineKeyboardButton("Отмена", callback_data="delete:no"),
    ]])
    await update.message.reply_text(
        "Удалить все твои отметки и отключить бот? Это действие необратимо.",
        reply_markup=keyboard,
    )


async def delete_confirm(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    assert query and query.data and update.effective_chat
    await query.answer()
    if query.data == "delete:no":
        await query.edit_message_text("Удаление отменено.")
        return
    with db() as conn:
        conn.execute("DELETE FROM checkins WHERE chat_id = ?", (update.effective_chat.id,))
        conn.execute("DELETE FROM participants WHERE chat_id = ?", (update.effective_chat.id,))
    await query.edit_message_text("Все данные удалены, напоминания отключены.")


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    context.user_data.clear()
    if update.message:
        await update.message.reply_text("Отметка отменена.")
    return ConversationHandler.END


def schedule_jobs(application: Application) -> None:
    for item in CHECKIN_TIMES:
        hour, minute = map(int, item.split(":"))
        application.job_queue.run_daily(
            scheduled_prompt,
            time=datetime.now(TZ).replace(hour=hour, minute=minute, second=0, microsecond=0).timetz(),
            name=f"checkin-{item}",
        )


def build_application(token: str) -> Application:
    application = Application.builder().token(token).build()
    conversation = ConversationHandler(
        entry_points=[
            CommandHandler("start", start),
            CommandHandler("checkin", begin_checkin),
            CallbackQueryHandler(begin_button, pattern=r"^begin:"),
        ],
        states={
            CONSENT: [CallbackQueryHandler(consent, pattern=r"^consent:")],
            MOOD: [CallbackQueryHandler(numeric_answer, pattern=r"^mood:")],
            ENERGY: [CallbackQueryHandler(numeric_answer, pattern=r"^energy:")],
            ANXIETY: [CallbackQueryHandler(numeric_answer, pattern=r"^anxiety:")],
            SLEEP: [MessageHandler(filters.TEXT & ~filters.COMMAND, sleep_answer), CommandHandler("skip", skip_sleep)],
            MEDS: [CallbackQueryHandler(meds_answer, pattern=r"^meds:")],
            SIDE_EFFECTS: [
                MessageHandler(filters.TEXT & ~filters.COMMAND, side_effects_answer),
                CommandHandler("skip", skip_side_effects),
            ],
            NOTE: [MessageHandler(filters.TEXT & ~filters.COMMAND, note_answer), CommandHandler("skip", skip_note)],
            SAFETY: [CallbackQueryHandler(safety_answer, pattern=r"^safety:")],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
    )
    application.add_handler(conversation)
    application.add_handler(CommandHandler("report", report))
    application.add_handler(CommandHandler("export", export))
    application.add_handler(CommandHandler("pause", pause))
    application.add_handler(CommandHandler("resume", resume))
    application.add_handler(CommandHandler("delete_my_data", delete_my_data))
    application.add_handler(CommandHandler("id", my_id))
    application.add_handler(CallbackQueryHandler(delete_confirm, pattern=r"^delete:"))
    return application


def main() -> None:
    token = os.getenv("BOT_TOKEN", "").strip()
    if not token:
        raise RuntimeError("Не задан BOT_TOKEN. Скопируйте .env.example в .env и укажите новый токен.")
    init_db()
    application = build_application(token)
    schedule_jobs(application)
    application.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
